import json
import io
import os
from openai import OpenAI
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Preset Warna & Tema Visual PPT
TEMPLATES = {
    "modern": {
        "bg_color": RGBColor(245, 247, 250),
        "title_color": RGBColor(20, 30, 55),
        "text_color": RGBColor(60, 64, 67),
    },
    "kreatif": {
        "bg_color": RGBColor(255, 250, 240),
        "title_color": RGBColor(180, 50, 20),
        "text_color": RGBColor(40, 40, 40),
    },
    "formal": {
        "bg_color": RGBColor(255, 255, 255),
        "title_color": RGBColor(10, 45, 90),
        "text_color": RGBColor(30, 30, 30),
    }
}

def generate_slide_content(topik, target_kelas="Kelas 10"):
    client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
    
    prompt = f"""
    Buat materi presentasi bimbingan klasikal untuk siswa {target_kelas} dengan topik: "{topik}".
    
    Output HARUS dalam format JSON dengan struktur:
    {{
        "slides": [
            {{
                "slide_title": "Judul Slide",
                "bullet_points": ["Poin 1", "Poin 2", "Poin 3"],
                "speaker_notes": "Catatan ringkas instruktur"
            }}
        ]
    }}
    Buat 4-6 slide yang interaktif.
    """
    
    response = client.chat.completions.create(
        model="gpt-4o-mini", 
        response_format={ "type": "json_object" },
        messages=[
            {"role": "system", "content": "Anda adalah desainer instruksional bimbingan konseling yang hanya mengeluarkan output JSON valid."},
            {"role": "user", "content": prompt}
        ]
    )
    
    data = json.loads(response.choices[0].message.content)
    return data.get("slides", [])

def build_pptx(slides_data, template_name="modern"):
    prs = Presentation()
    theme = TEMPLATES.get(template_name, TEMPLATES["modern"])
    blank_layout = prs.slide_layouts[6] 

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        
        # 1. Background Warna
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = theme["bg_color"]

        # 2. Judul Slide
        tx_box_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.4), Inches(1.2))
        tf_title = tx_box_title.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = data.get("slide_title", "")
        p_title.font.bold = True
        p_title.font.size = Pt(28)
        p_title.font.color.rgb = theme["title_color"]

        # 3. Poin Materi (Bullet Points)
        tx_box_body = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(8.4), Inches(4.5))
        tf_body = tx_box_body.text_frame
        tf_body.word_wrap = True
        
        for point in data.get("bullet_points", []):
            p = tf_body.add_paragraph()
            p.text = f"•  {point}"
            p.font.size = Pt(18)
            p.font.color.rgb = theme["text_color"]
            p.space_after = Pt(14)

    output_stream = io.BytesIO()
    prs.save(output_stream)
    output_stream.seek(0)
    return output_stream